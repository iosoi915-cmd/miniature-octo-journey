# -*- coding: utf-8 -*-
"""
홍성 창의센터 회원 수강 현황 분석 프로그램

구글 시트(또는 구글 설문 응답 시트)에서 내려받은 수강 기록을 읽어
회원별 단계 이수 현황, 재수강 현황, 현재 수강 현황을 엑셀 리포트로 만듭니다.

사용법
    python member_report.py                       # 실행 후 파일 경로/링크 입력
    python member_report.py 수강기록.csv           # CSV 파일
    python member_report.py 수강기록.xlsx          # 엑셀 파일
    python member_report.py "https://docs.google.com/spreadsheets/d/...."   # 구글 시트 링크
    python member_report.py 수강기록.csv --member 홍길동   # 특정 회원 조회
    python member_report.py --sample               # 샘플 데이터로 체험

필요한 패키지:  pip install pandas openpyxl
"""
import argparse
import random
import re
import sys
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ============================================================
# 설정 (센터 상황에 맞게 수정하세요)
# ============================================================

# 구글 시트 공유 링크 ('링크가 있는 모든 사용자 - 뷰어'로 공유 필요).
# 여기에 넣어두면 실행할 때마다 입력하지 않아도 됩니다.
GOOGLE_SHEET_URL = ""

# 단계 순서 (앞에서부터 1단계)
STAGE_COURSES = ["씨앗반", "새싹반", "나무반", "열매반"]

# 단계와 별개로 운영되는 특별반
SPECIAL_COURSES = ["그래비트랙스반"]

# 반 이름에 이 글자가 들어 있으면 해당 반으로 인식 ("씨앗반 3기", "새싹 A반" 등)
COURSE_KEYWORDS = {
    "씨앗": "씨앗반",
    "새싹": "새싹반",
    "나무": "나무반",
    "열매": "열매반",
    "그래비": "그래비트랙스반",
    "gravi": "그래비트랙스반",
}

# 시트의 열 이름 후보. 시트 열 이름이 다르면 여기에 추가하세요.
COLUMN_ALIASES = {
    "member_id": ["회원번호", "회원ID", "회원아이디"],
    "name": ["이름", "성명", "회원명", "학생명", "학생이름", "아동이름"],
    "phone": ["연락처", "전화번호", "휴대폰", "핸드폰", "보호자연락처"],
    "course": ["수강반", "반", "과정", "수강과정", "수업", "강좌", "프로그램"],
    "term": ["기수", "학기", "차수", "회차"],
    "start": ["시작일", "개강일", "수강시작일", "타임스탬프"],
    "end": ["종료일", "종강일", "수강종료일"],
    "status": ["이수여부", "수료여부", "상태", "수강상태"],
}

# 상태 값: 수료 / 수강중 / 수강예정 / 미수료 / 중도포기 / 미확인
DONE, CURRENT, PLANNED, FAILED, DROPPED, UNKNOWN = "수료", "수강중", "수강예정", "미수료", "중도포기", "미확인"


# ============================================================
# 데이터 읽기
# ============================================================

def to_csv_url(url):
    """구글 시트 공유 링크를 CSV 다운로드 링크로 바꾼다."""
    m = re.search(r"/spreadsheets/d/([\w-]+)", url)
    if not m:
        return url
    out = f"https://docs.google.com/spreadsheets/d/{m[1]}/export?format=csv"
    gid = re.search(r"[#&?]gid=(\d+)", url)
    if gid:
        out += f"&gid={gid[1]}"
    return out


def load_table(src, sheet=None):
    if re.match(r"https?://", src):
        return pd.read_csv(to_csv_url(src), dtype=str)
    path = Path(src)
    if not path.exists():
        raise SystemExit(f"[오류] 파일을 찾을 수 없습니다: {path}")
    if path.suffix.lower() in (".xlsx", ".xlsm", ".xls"):
        return pd.read_excel(path, sheet_name=sheet or 0, dtype=str)
    for enc in ("utf-8-sig", "cp949"):
        try:
            return pd.read_csv(path, dtype=str, encoding=enc)
        except UnicodeDecodeError:
            continue
    raise SystemExit("[오류] CSV 파일 인코딩을 읽을 수 없습니다. UTF-8로 저장해 주세요.")


def squash(s):
    return re.sub(r"\s+", "", str(s))


def map_columns(df):
    """시트 열 이름을 내부 필드 이름에 연결한다 (정확히 일치 우선, 그다음 포함)."""
    norm = {c: squash(c) for c in df.columns}
    found, used = {}, set()
    for field, aliases in COLUMN_ALIASES.items():
        al = [squash(a) for a in aliases]
        col = next((c for a in al for c in df.columns if c not in used and norm[c] == a), None)
        if col is None:
            col = next((c for a in al for c in df.columns if c not in used and a in norm[c]), None)
        if col is not None:
            found[field] = col
            used.add(col)
    missing = [f for f in ("name", "course") if f not in found]
    if missing:
        raise SystemExit(
            "[오류] 필수 열(이름, 수강반)을 찾지 못했습니다.\n"
            f"  시트의 열: {list(df.columns)}\n"
            "  → COLUMN_ALIASES 설정에 열 이름을 추가하거나, 구글 시트 공유 권한을 확인하세요."
        )
    return found


def parse_date(v):
    if v is None:
        return pd.NaT
    s = str(v).strip()
    if not s or s.lower() in ("nan", "nat", "none"):
        return pd.NaT
    # "2025-03-01", "2025. 3. 1 오후 2:30:00", "2025/3/1" 등
    m = re.search(r"(\d{4})\D{1,3}(\d{1,2})\D{1,3}(\d{1,2})", s)
    if m:
        try:
            return pd.Timestamp(int(m[1]), int(m[2]), int(m[3]))
        except ValueError:
            return pd.NaT
    return pd.to_datetime(s, errors="coerce")


def normalize_course(v):
    s = squash(v).lower()
    for key, course in COURSE_KEYWORDS.items():
        if key in s:
            return course
    return str(v).strip()


def decide_status(raw, start, end, today):
    s = squash(raw).lower()
    if s and s != "nan":
        if any(k in s for k in ("미수료", "미이수", "불합격")):
            return FAILED
        if any(k in s for k in ("포기", "중도", "취소", "환불", "중단")):
            return DROPPED
        if any(k in s for k in ("예정", "대기")):
            return PLANNED
        if any(k in s for k in ("수강중", "진행", "재원")):
            return CURRENT
        if any(k in s for k in ("수료", "이수", "완료")) or s in ("o", "y", "yes", "true", "v"):
            return DONE
    # 상태가 비어 있으면 날짜로 판단
    if pd.notna(end) and end < today:
        return DONE
    if pd.notna(start) and start > today:
        return PLANNED
    if pd.notna(start) or pd.notna(end):
        return CURRENT
    return UNKNOWN


def clean(raw, today):
    cols = map_columns(raw)

    def get(field):
        if field not in cols:
            return pd.Series([""] * len(raw), index=raw.index)
        return raw[cols[field]].fillna("").astype(str).str.strip()

    df = pd.DataFrame({
        "member_id": get("member_id"),
        "name": get("name"),
        "phone": get("phone"),
        "course_raw": get("course"),
        "term": get("term"),
        "start": get("start").map(parse_date),
        "end": get("end").map(parse_date),
        "status_raw": get("status"),
    })
    df = df[(df["name"] != "") & (df["course_raw"] != "")]

    # 구글 설문 체크박스처럼 "씨앗반, 그래비트랙스반" 이 한 칸에 있으면 나눈다
    df = df.assign(course_raw=df["course_raw"].str.split(r"[,/]")).explode("course_raw")
    df["course_raw"] = df["course_raw"].str.strip()
    df = df[df["course_raw"] != ""].copy()
    if df.empty:
        raise SystemExit("[오류] 분석할 수강 기록이 없습니다.")

    df["course"] = df["course_raw"].map(normalize_course)
    df["status"] = [decide_status(r, s, e, today)
                    for r, s, e in zip(df["status_raw"], df["start"], df["end"])]

    # 회원 구분: 회원번호 > 이름+연락처 뒷4자리 > 이름 (동명이인 구분용)
    def member_key(r):
        if r.member_id:
            return f"ID:{r.member_id}"
        digits = re.sub(r"\D", "", r.phone)
        return f"{r.name}_{digits[-4:]}" if digits else r.name

    df["key"] = [member_key(r) for r in df.itertuples()]
    return df.sort_values(["name", "key", "start"], na_position="last").reset_index(drop=True)


# ============================================================
# 분석
# ============================================================

def course_order(df):
    known = STAGE_COURSES + SPECIAL_COURSES
    return known + sorted(set(df["course"]) - set(known))


def fmt_date(d):
    return d.strftime("%Y-%m-%d") if pd.notna(d) else ""


def rec_label(r):
    when = r.term or fmt_date(r.start) or "날짜없음"
    return f"{when}({r.status})"


def taken(df):
    """실제로 수강한 기록 (수강예정 제외)."""
    return df[df["status"] != PLANNED]


def member_summary(df):
    rows = []
    for key, g in df.groupby("key", sort=False):
        done = set(g.loc[g["status"] == DONE, "course"])
        idx = max((i for i, c in enumerate(STAGE_COURSES) if c in done), default=-1)
        current = [c for c in course_order(df) if ((g["course"] == c) & (g["status"] == CURRENT)).any()]
        planned = [c for c in course_order(df) if ((g["course"] == c) & (g["status"] == PLANNED)).any()]

        if idx == len(STAGE_COURSES) - 1:
            next_stage = "전 단계 이수 완료"
        else:
            next_stage = STAGE_COURSES[idx + 1]
            if next_stage in current:
                next_stage += " (수강중)"
            elif next_stage in planned:
                next_stage += " (등록함)"

        counts = taken(g)["course"].value_counts()
        repeats = [f"{c}({counts[c]}회)" for c in course_order(df) if counts.get(c, 0) >= 2]
        skipped = [c for c in STAGE_COURSES[:max(idx, 0)] if c not in done]

        row = {
            "회원": g["name"].iloc[0],
            "연락처": next((p for p in g["phone"] if p), ""),
            "최고 이수 단계": STAGE_COURSES[idx] if idx >= 0 else "-",
            "이수 단계 수": f"{idx + 1}/{len(STAGE_COURSES)}",
            "다음 단계": next_stage,
            "현재 수강중": ", ".join(current),
            "수강 예정": ", ".join(planned),
            "총 수강 횟수": int(len(taken(g))),
            "재수강 여부": "재수강" if repeats else "",
            "재수강 과정": ", ".join(repeats),
            "건너뛴 단계": ", ".join(skipped),
        }
        for c in SPECIAL_COURSES:
            row[f"{c} 수료"] = int(((g["course"] == c) & (g["status"] == DONE)).sum())
        row["첫 수강일"] = fmt_date(g["start"].min())
        row["최근 수강일"] = fmt_date(g["start"].max())
        rows.append(row)
    return pd.DataFrame(rows)


def stage_matrix(df):
    courses = course_order(df)
    rows = []
    for key, g in df.groupby("key", sort=False):
        row = {"회원": g["name"].iloc[0], "연락처": next((p for p in g["phone"] if p), "")}
        for c in courses:
            st = g.loc[g["course"] == c, "status"]
            parts = []
            if (n := (st == DONE).sum()):
                parts.append(f"수료 {n}회")
            if (st == CURRENT).any():
                parts.append("수강중")
            if (st == PLANNED).any():
                parts.append("예정")
            if (n := st.isin([FAILED, DROPPED]).sum()):
                parts.append(f"미수료/포기 {n}회")
            if (n := (st == UNKNOWN).sum()):
                parts.append(f"기록 {n}회")
            row[c] = " / ".join(parts)
        rows.append(row)
    return pd.DataFrame(rows)


def repeat_list(df):
    rows = []
    t = taken(df)
    for (key, course), g in t.groupby(["key", "course"], sort=False):
        if len(g) < 2:
            continue
        rows.append({
            "회원": g["name"].iloc[0],
            "연락처": next((p for p in g["phone"] if p), ""),
            "과정": course,
            "수강 횟수": len(g),
            "수료 횟수": int((g["status"] == DONE).sum()),
            "수강 이력": ", ".join(rec_label(r) for r in g.itertuples()),
        })
    cols = ["회원", "연락처", "과정", "수강 횟수", "수료 횟수", "수강 이력"]
    out = pd.DataFrame(rows, columns=cols)
    order = {c: i for i, c in enumerate(course_order(df))}
    return out.sort_values(["과정", "수강 횟수"], key=lambda s: s.map(order) if s.name == "과정" else -s)


def current_list(df):
    cur = df[df["status"].isin([CURRENT, PLANNED])]
    rows = []
    for r in cur.itertuples():
        done = set(df.loc[(df["key"] == r.key) & (df["status"] == DONE), "course"])
        prev = [c for c in STAGE_COURSES if c in done]
        n_before = int(((df["key"] == r.key) & (df["course"] == r.course) & (df["status"] != PLANNED)).sum())
        rows.append({
            "상태": r.status,
            "과정": r.course,
            "회원": r.name,
            "연락처": r.phone,
            "기수": r.term,
            "시작일": fmt_date(r.start),
            "종료일": fmt_date(r.end),
            "이 과정 수강 횟수": n_before,
            "이수한 단계": ", ".join(prev),
        })
    cols = ["상태", "과정", "회원", "연락처", "기수", "시작일", "종료일", "이 과정 수강 횟수", "이수한 단계"]
    out = pd.DataFrame(rows, columns=cols)
    order = {c: i for i, c in enumerate(course_order(df))}
    return out.sort_values(["상태", "과정", "회원"], key=lambda s: s.map(order) if s.name == "과정" else s)


def course_stats(df):
    rows = []
    for c in course_order(df):
        g = df[df["course"] == c]
        t = taken(g)
        per_member = t.groupby("key").size()
        members = len(per_member)
        completed = set(g.loc[g["status"] == DONE, "key"])
        repeaters = int((per_member >= 2).sum())
        row = {
            "과정": c,
            "구분": f"{STAGE_COURSES.index(c) + 1}단계" if c in STAGE_COURSES else "특별반",
            "총 수강 건수": len(t),
            "수강 회원 수": members,
            "수료 회원 수": len(completed),
            "재수강 회원 수": repeaters,
            "재수강률": f"{repeaters / members:.0%}" if members else "-",
            "현재 수강중": int(g.loc[g["status"] == CURRENT, "key"].nunique()),
            "수강 예정": int(g.loc[g["status"] == PLANNED, "key"].nunique()),
            "미수료/포기 건수": int(g["status"].isin([FAILED, DROPPED]).sum()),
            "다음 단계 진급": "-",
            "진급률": "-",
        }
        if c in STAGE_COURSES[:-1]:
            nxt = STAGE_COURSES[STAGE_COURSES.index(c) + 1]
            promoted = completed & set(df.loc[df["course"] == nxt, "key"])
            row["다음 단계 진급"] = f"{len(promoted)}명 → {nxt}"
            row["진급률"] = f"{len(promoted) / len(completed):.0%}" if completed else "-"
        rows.append(row)
    return pd.DataFrame(rows)


def records_sheet(df):
    return pd.DataFrame({
        "회원": df["name"],
        "연락처": df["phone"],
        "과정": df["course"],
        "원래 입력값": df["course_raw"],
        "기수": df["term"],
        "시작일": df["start"].map(fmt_date),
        "종료일": df["end"].map(fmt_date),
        "상태": df["status"],
        "회원 구분키": df["key"],
    })


# ============================================================
# 엑셀 저장
# ============================================================

HEADER_FILL = PatternFill("solid", fgColor="2F5496")
FILL_DONE = PatternFill("solid", fgColor="C6EFCE")
FILL_CURRENT = PatternFill("solid", fgColor="FFEB9C")
FILL_BAD = PatternFill("solid", fgColor="FFC7CE")


def text_width(v):
    s = "" if v is None else str(v)
    return sum(2 if ord(ch) > 127 else 1 for ch in s)


def style_sheet(ws, color_cells=False):
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        width = max(text_width(c.value) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(width + 2, 8), 60)
    if color_cells:
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                v = str(cell.value or "")
                if "수강중" in v or v.endswith("예정"):
                    cell.fill = FILL_CURRENT
                elif v.startswith("수료"):
                    cell.fill = FILL_DONE
                elif v in (FAILED, DROPPED) or "포기" in v:
                    cell.fill = FILL_BAD


def write_report(df, out_path):
    sheets = [
        ("회원별 현황", member_summary(df), True),
        ("단계별 이수표", stage_matrix(df), True),
        ("재수강 현황", repeat_list(df), False),
        ("현재 수강 현황", current_list(df), True),
        ("반별 통계", course_stats(df), False),
        ("전체 기록", records_sheet(df), True),
    ]
    try:
        with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
            for name, table, color in sheets:
                table.to_excel(writer, sheet_name=name, index=False)
                style_sheet(writer.sheets[name], color)
    except PermissionError:
        raise SystemExit(f"[오류] {out_path} 파일이 열려 있습니다. 엑셀을 닫고 다시 실행하세요.")
    return dict((n, t) for n, t, _ in sheets)


# ============================================================
# 화면 출력
# ============================================================

def print_summary(df, tables):
    summary = tables["회원별 현황"]
    print("\n" + "=" * 50)
    print(" 홍성 창의센터 수강 현황 요약")
    print("=" * 50)
    print(f" 전체 회원 수      : {len(summary)}명")
    print(f" 전체 수강 건수    : {len(taken(df))}건")
    print(f" 현재 수강중 회원  : {df.loc[df['status'] == CURRENT, 'key'].nunique()}명")
    print(f" 재수강 회원       : {(summary['재수강 여부'] == '재수강').sum()}명")
    print("-" * 50)
    print(" [단계별 이수 회원 수]")
    for c in STAGE_COURSES + SPECIAL_COURSES:
        n = df.loc[(df["course"] == c) & (df["status"] == DONE), "key"].nunique()
        cur = df.loc[(df["course"] == c) & (df["status"] == CURRENT), "key"].nunique()
        print(f"   {c:<10} 이수 {n:>3}명   수강중 {cur:>3}명")
    print("-" * 50)
    print(" [최고 이수 단계별 회원 분포]")
    for stage, n in summary["최고 이수 단계"].value_counts().reindex(["-"] + STAGE_COURSES, fill_value=0).items():
        label = "이수 단계 없음" if stage == "-" else stage
        print(f"   {label:<12} {n:>3}명")
    unknown = (df["status"] == UNKNOWN).sum()
    if unknown:
        print(f"\n ※ 상태/날짜가 없어 판단하지 못한 기록 {unknown}건 ('미확인')")
    print("=" * 50)


def print_member(df, name):
    hit = df[df["name"].str.contains(name, regex=False)]
    if hit.empty:
        print(f"\n'{name}' 회원을 찾을 수 없습니다.")
        return
    for key, g in hit.groupby("key", sort=False):
        phone = next((p for p in g["phone"] if p), "")
        print(f"\n■ {g['name'].iloc[0]}  {phone}")
        for r in g.itertuples():
            period = f"{fmt_date(r.start)} ~ {fmt_date(r.end)}" if pd.notna(r.start) else ""
            print(f"   - {r.course:<10} {r.term:<12} {period:<25} {r.status}")
        done = set(g.loc[g["status"] == DONE, "course"])
        print(f"   이수 단계: {', '.join(c for c in STAGE_COURSES if c in done) or '없음'}")


# ============================================================
# 샘플 데이터
# ============================================================

def make_sample(path, today, n=60, seed=7):
    rnd = random.Random(seed)
    last = "김이박최정강조윤장임한오서신권황안송류홍"
    first = ["민준", "서연", "도윤", "하윤", "시우", "지우", "주원", "서윤", "예준", "하은",
             "지호", "수아", "준우", "지유", "현우", "채원", "건우", "다은", "우진", "소율"]
    terms, t = [], pd.Timestamp(2024, 3, 1)
    while t <= today + pd.DateOffset(months=3):
        terms.append(t)
        t += pd.DateOffset(months=3)

    rows = []
    for _ in range(n):
        name = rnd.choice(last) + rnd.choice(first)
        phone = f"010-{rnd.randint(1000, 9999)}-{rnd.randint(1000, 9999)}"
        stage = 0
        grav = set(rnd.sample(range(len(terms)), k=rnd.choice([0, 0, 1, 2])))
        for k in range(rnd.randrange(len(terms)), len(terms)):
            start = terms[k]
            end = start + pd.DateOffset(months=3) - pd.Timedelta(days=1)
            label = f"{start.year}년 {start.month // 3}기"
            live = PLANNED if start > today else CURRENT if end >= today else None

            def add(course, status):
                rows.append([name, phone, course, label, fmt_date(start), fmt_date(end), status])

            if k in grav:
                add("그래비트랙스반", live or DONE)
            if stage >= len(STAGE_COURSES):
                continue
            if live:
                add(STAGE_COURSES[stage], live)
                break
            r = rnd.random()
            if r < 0.68:
                add(STAGE_COURSES[stage], DONE)
                stage += 1
            elif r < 0.83:
                add(STAGE_COURSES[stage], DONE)   # 수료 후 같은 단계 재수강
            elif r < 0.93:
                add(STAGE_COURSES[stage], FAILED)
            else:
                add(STAGE_COURSES[stage], DROPPED)
                break
            if rnd.random() < 0.1:              # 휴원
                break

    pd.DataFrame(rows, columns=["이름", "연락처", "수강반", "기수", "시작일", "종료일", "이수여부"]) \
        .to_csv(path, index=False, encoding="utf-8-sig")
    return path


# ============================================================
# 실행
# ============================================================

def main():
    ap = argparse.ArgumentParser(description="홍성 창의센터 회원 수강 현황 분석")
    ap.add_argument("source", nargs="?", help="CSV/엑셀 파일 경로 또는 구글 시트 링크")
    ap.add_argument("--sheet", help="엑셀 파일의 시트 이름 (기본: 첫 번째 시트)")
    ap.add_argument("--out", help="결과 엑셀 파일 이름")
    ap.add_argument("--member", help="특정 회원의 수강 이력 조회 (이름 일부만 입력해도 됨)")
    ap.add_argument("--today", help="기준 날짜 YYYY-MM-DD (기본: 오늘)")
    ap.add_argument("--sample", action="store_true", help="샘플 데이터를 만들어 분석")
    args = ap.parse_args()
    interactive = len(sys.argv) == 1

    today = pd.Timestamp(args.today) if args.today else pd.Timestamp(date.today())

    if args.sample:
        sample = make_sample(Path("샘플_수강기록.csv"), today)
        print(f"샘플 데이터를 만들었습니다: {sample.resolve()}")
        args.source = args.source or str(sample)

    src = args.source or GOOGLE_SHEET_URL
    if not src:
        src = input("수강 기록 파일 경로 또는 구글 시트 링크를 입력하세요:\n> ").strip().strip('"')
        if not src:
            raise SystemExit("입력이 없어 종료합니다.")

    try:
        raw = load_table(src, args.sheet)
    except SystemExit:
        raise
    except Exception as e:
        raise SystemExit(f"[오류] 데이터를 읽지 못했습니다: {e}\n"
                         "  구글 시트라면 공유 설정이 '링크가 있는 모든 사용자'인지 확인하세요.")

    df = clean(raw, today)

    if args.out:
        out = Path(args.out)
    else:
        base = Path.cwd() if re.match(r"https?://", src) else Path(src).resolve().parent
        out = base / f"수강현황_리포트_{today:%Y%m%d}.xlsx"

    tables = write_report(df, out)
    print_summary(df, tables)
    print(f"\n엑셀 리포트를 저장했습니다: {out.resolve()}")

    if args.member:
        print_member(df, args.member)

    if interactive:
        while True:
            name = input("\n조회할 회원 이름 (엔터 = 종료): ").strip()
            if not name:
                break
            print_member(df, name)


if __name__ == "__main__":
    try:
        main()
    except SystemExit as e:
        if e.code not in (None, 0):
            print(e.code if isinstance(e.code, str) else "", file=sys.stderr)
            if len(sys.argv) == 1:
                input("\n엔터를 누르면 종료합니다.")
        sys.exit(1 if e.code not in (None, 0) else 0)
